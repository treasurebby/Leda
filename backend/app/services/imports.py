"""Spreadsheet parsing for product and retailer imports.

Mirrors frontend/src/onboarding/imports.ts: same header aliases, limits and error wording, so a file that
passes client-side validation passes here too, and vice versa.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.core.config import settings
from app.core.security import normalise_phone

EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]{2,}$")


class ImportError_(Exception):
    """Raised for whole-file problems (wrong type, missing columns) or when rows fail validation."""


def _header_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def _column(headers: list[str], aliases: list[str]) -> int:
    keys = {_header_key(a) for a in aliases}
    for i, h in enumerate(headers):
        if _header_key(h) in keys:
            return i
    return -1


def _cell(row: list[str], index: int) -> str:
    return row[index] if 0 <= index < len(row) else ""


def _numeric(value: str, fallback: Decimal | None = None) -> Decimal | None:
    if value == "":
        return fallback
    cleaned = re.sub(r"NGN|₦|,", "", value, flags=re.I).strip()
    try:
        return Decimal(cleaned) if cleaned else None
    except InvalidOperation:
        return None


def read_spreadsheet(file_name: str, content: bytes) -> list[list[str]]:
    if not content:
        raise ImportError_("This file is empty. Add a header row and at least one record.")
    if len(content) > settings.import_max_bytes:
        raise ImportError_("This file is a little too large. Please use a file under 5 MB.")
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if ext not in {"csv", "xlsx"}:
        raise ImportError_(
            "Please choose a CSV or Excel (.xlsx) file. Older .xls files can be saved as .xlsx in Excel."
        )

    rows: list[list[str]]
    if ext == "csv":
        text = content.decode("utf-8-sig", errors="replace")
        try:
            rows = list(csv.reader(io.StringIO(text)))
        except csv.Error as exc:
            raise ImportError_(
                "Some quotation marks in this CSV do not match. Please check the file and try again."
            ) from exc
    else:
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
        except Exception as exc:
            raise ImportError_(
                "We couldn't read this workbook. Please check it isn't password protected, or save it as CSV."
            ) from exc

    rows = [[c.strip() for c in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if len(rows) < 2:
        raise ImportError_("Add column headers and at least one record to your file.")
    if len(rows) > settings.import_max_rows + 1:
        raise ImportError_("Please import up to 5,000 records at a time.")
    return rows


def _raise_row_errors(errors: list[str]) -> None:
    if errors:
        head = " ".join(errors[:3])
        tail = f" Plus {len(errors) - 3} more issues." if len(errors) > 3 else ""
        raise ImportError_(head + tail)


@dataclass
class ProductRow:
    name: str
    sku: str
    price: Decimal
    stock: int
    unit: str | None = None


@dataclass
class RetailerRow:
    name: str
    email: str | None
    phone: str | None
    market: str | None = None
    aliases: list[str] = field(default_factory=list)


def parse_products(rows: list[list[str]]) -> list[ProductRow]:
    headers, *data = rows
    name_col = _column(headers, ["product", "product name", "name", "item", "item name"])
    price_col = _column(headers, ["price", "price ngn", "unit price", "unit price ngn", "selling price", "amount"])
    sku_col = _column(headers, ["sku", "product code", "item code", "code"])
    stock_col = _column(
        headers, ["stock", "quantity", "qty", "stock quantity", "stock level", "quantity available", "stock on hand"]
    )
    unit_col = _column(headers, ["unit", "pack", "pack size", "unit size"])
    if name_col < 0 or price_col < 0:
        raise ImportError_(
            'Your file needs "Product name" and "Price" columns. SKU and Stock are optional. '
            "Our template has the right headings."
        )

    errors: list[str] = []
    seen: set[str] = set()
    products: list[ProductRow] = []
    for index, row in enumerate(data):
        line = index + 2
        name = _cell(row, name_col)
        price = _numeric(_cell(row, price_col))
        stock = _numeric(_cell(row, stock_col), Decimal(0))
        sku = (_cell(row, sku_col) or f"LEDA{index + 1:04d}").upper()
        if not name:
            errors.append(f"Row {line}: product name is missing.")
        if price is None or price < 0:
            errors.append(f"Row {line}: enter a valid, nonnegative price.")
        if stock is None or stock < 0 or stock != stock.to_integral_value():
            errors.append(f"Row {line}: stock must be a nonnegative whole number.")
        if sku in seen:
            errors.append(f"Row {line}: SKU {sku} is repeated.")
        seen.add(sku)
        products.append(
            ProductRow(
                name=name,
                sku=sku,
                price=price or Decimal(0),
                stock=int(stock or 0),
                unit=_cell(row, unit_col) or None,
            )
        )
    _raise_row_errors(errors)
    return products


def parse_retailers(rows: list[list[str]]) -> list[RetailerRow]:
    headers, *data = rows
    name_col = _column(headers, ["retailer name", "retailer", "business name", "name", "customer", "customer name"])
    email_col = _column(headers, ["email", "email address", "retailer email"])
    phone_col = _column(
        headers,
        [
            "phone",
            "phone number",
            "phone no",
            "mobile",
            "mobile number",
            "whatsapp",
            "whatsapp number",
            "telephone",
            "contact number",
        ],
    )
    market_col = _column(headers, ["market", "location", "area", "address"])
    if name_col < 0 or (email_col < 0 and phone_col < 0):
        raise ImportError_(
            'Your file needs a "Retailer name" column and at least one contact column: "Email" or "Phone".'
        )

    errors: list[str] = []
    contacts: set[str] = set()
    retailers: list[RetailerRow] = []
    for index, row in enumerate(data):
        line = index + 2
        name = _cell(row, name_col)
        email = _cell(row, email_col).lower()
        raw_phone = _cell(row, phone_col)
        phone = normalise_phone(raw_phone) if raw_phone else None
        if not name:
            errors.append(f"Row {line}: retailer name is missing.")
        if not email and not raw_phone:
            errors.append(f"Row {line}: add an email or phone number.")
        if email and not EMAIL_RE.match(email):
            errors.append(f"Row {line}: the email is not valid.")
        if raw_phone and not phone:
            errors.append(f"Row {line}: use a valid Nigerian phone number.")
        key = email or phone
        if key and key in contacts:
            errors.append(f"Row {line}: this contact appears more than once.")
        if key:
            contacts.add(key)
        retailers.append(
            RetailerRow(name=name, email=email or None, phone=phone, market=_cell(row, market_col) or None)
        )
    _raise_row_errors(errors)
    return retailers
